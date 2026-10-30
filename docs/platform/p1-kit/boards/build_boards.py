import sys
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.chart import LineChart, Reference, Series
OUT=sys.argv[1]; F="Arial"
YEL=PatternFill("solid",fgColor="FFF2CC"); HEAD=PatternFill("solid",fgColor="DDE6F5"); KEY=PatternFill("solid",fgColor="E2F0D9")
thin=Side(style="thin",color="BBBBBB"); BOX=Border(left=thin,right=thin,top=thin,bottom=thin)
def inp(c,v,fmt=None):
    c.value=v; c.fill=YEL; c.font=Font(name=F,color="0000FF"); c.border=BOX
    if fmt: c.number_format=fmt
def calc(c,v,fmt=None,key=False):
    c.value=v; c.border=BOX; c.font=Font(name=F,bold=key)
    if key: c.fill=KEY
    if fmt: c.number_format=fmt
def lab(ws,r,col,text,bold=False):
    c=ws.cell(r,col,text); c.font=Font(name=F,bold=bold); c.border=BOX; c.alignment=Alignment(wrap_text=True,vertical="center"); return c
def header(ws,row,cols):
    for i,h in enumerate(cols,1):
        c=ws.cell(row,i,h); c.font=Font(name=F,bold=True); c.fill=HEAD; c.border=BOX; c.alignment=Alignment(wrap_text=True,vertical="center")
    ws.row_dimensions[row].height=32
def usage(wb,title,lines):
    ws=wb.active; ws.title="使い方"
    ws["A1"]=title; ws["A1"].font=Font(name=F,bold=True,size=14)
    for i,l in enumerate(lines,3): ws.cell(i,1,l).font=Font(name=F)
    r=len(lines)+4
    ws.cell(r,1,"凡例").font=Font(name=F,bold=True)
    c=ws.cell(r+1,1,"黄色のセル（青字）＝あなたが入力する欄"); c.fill=YEL; c.font=Font(name=F,color="0000FF")
    c=ws.cell(r+2,1,"緑のセル＝自動で計算される「現在地」と「これからの予定」"); c.fill=KEY; c.font=Font(name=F)
    ws.cell(r+3,1,"その他の白いセル＝自動計算（触らない）").font=Font(name=F)
    ws.cell(r+4,1,"入っている数字は架空の記入例です。上書きしてお使いください。").font=Font(name=F)
    ws.cell(r+6,1,"このシートは、入力した数字から機械的に計算した目安です。結果を保証するものではありません。氏名・口座番号・ID・パスワード・学校名などは書かないでください。").font=Font(name=F)
    ws.column_dimensions["A"].width=110
def chart(ws,src,title,ytitle,first,last,cols,anchor,cat_col=1):
    ch=LineChart(); ch.title=title; ch.y_axis.title=ytitle; ch.height=9; ch.width=20
    for col in cols: ch.add_data(Reference(src,min_col=col,min_row=first-1,max_row=last),titles_from_data=True)
    ch.set_categories(Reference(src,min_col=cat_col,min_row=first,max_row=last)); ws.add_chart(ch,anchor)

MROW=27  # スマホシートで毎月の数字を入れる欄の開始行
def mobile(wb,title):
    m=wb.create_sheet("スマホ",0); wb.active=0
    m["A1"]=title; m["A1"].font=Font(name=F,bold=True,size=13); m["A1"].alignment=Alignment(wrap_text=True)
    m.column_dimensions["A"].width=20; m.column_dimensions["B"].width=18
    return m
def mres(m,r,label,formula,fmt=None,key=True):
    c=m.cell(r,1,label); c.font=Font(name=F); c.border=BOX; c.alignment=Alignment(wrap_text=True,vertical="center")
    calc(m.cell(r,2),formula,fmt,key); m.cell(r,2).alignment=Alignment(wrap_text=True,vertical="center")
def plan_chart(src,mob,n,title,ytitle,small):
    """当初の計画（記録B）・実績（スマホの入力欄＝空欄はそのまま空欄）・更新後の予定（記録D）"""
    ch=LineChart(); ch.title=title; ch.y_axis.title=ytitle
    ch.height,ch.width=(7,9.5) if small else (9,20)
    ch.legend.position="b"
    for ref,name in [(Reference(src,min_col=2,min_row=2,max_row=n+1),"当初の計画"),
                     (Reference(mob,min_col=2,min_row=MROW,max_row=MROW+n-1),"実績"),
                     (Reference(src,min_col=4,min_row=2,max_row=n+1),"更新後の予定")]:
        sr=Series(ref,title=name); ch.series.append(sr)
    ch.set_categories(Reference(src,min_col=1,min_row=2,max_row=n+1))
    return ch
def mchart(m,src,title,ytitle,first,last,cols,anchor="A12"):
    m.add_chart(plan_chart(src,m,last-1,title,ytitle,True),anchor)
def fix_heights(m):
    for r in range(3,10): m.row_dimensions[r].height=44
    m.merge_cells("A10:B10"); m.row_dimensions[10].height=56; m["A10"].alignment=Alignment(wrap_text=True,vertical="top")

def linear(title,usage_lines,unit,start,goal,n,period_word,actuals,metric_name,fmt="0.0",extra=None):
    """開始値→目標値をn期間で。実績から現在地・必要ペース・見込み・更新後の予定を計算。"""
    wb=Workbook(); usage(wb,title,usage_lines)
    p=wb.create_sheet("前提"); p["A1"]="前提（初回の計画書から転記）"; p["A1"].font=Font(name=F,bold=True)
    lab(p,2,1,f"{metric_name}：最初の値（{unit}）"); inp(p.cell(2,2),start,fmt)
    lab(p,3,1,f"{metric_name}：目標の値（{unit}）"); inp(p.cell(3,2),goal,fmt)
    lab(p,4,1,f"期限までの{period_word}数"); inp(p.cell(4,2),n,"0")
    p.column_dimensions["A"].width=40; p.column_dimensions["B"].width=14
    m=mobile(wb,title.replace("現在地ボード","スマホで見る"))
    d=wb.create_sheet("記録")
    header(d,1,[period_word,f"当初の計画（{unit}）",f"実績（{unit}）※スマホのシートから自動",f"更新後の予定（{unit}）",f"今のペースが続いた場合（{unit}）"])
    last=n+1
    for k in range(1,n+1):
        r=k+1
        calc(d.cell(r,1),k,"0")
        calc(d.cell(r,2),f"=前提!$B$2+(前提!$B$3-前提!$B$2)*A{r}/前提!$B$4",fmt)
        mr=MROW+k-1
        lab(m,mr,1,f"{k}か月目"); inp(m.cell(mr,2),actuals[k-1] if k-1<len(actuals) else None,fmt)
        calc(d.cell(r,3),f'=IF(スマホ!B{mr}="","",スマホ!B{mr})',fmt)
        calc(d.cell(r,4),f'=IFERROR(IF(OR(ボード!$B$4="",ボード!$B$10="",A{r}<=ボード!$B$4),"",ボード!$B$5+ボード!$B$10*(A{r}-ボード!$B$4)),"")',fmt)
        calc(d.cell(r,5),f'=IFERROR(IF(OR(ボード!$B$4="",A{r}<=ボード!$B$4,ボード!$B$11=""),"",ボード!$B$5+ボード!$B$11*(A{r}-ボード!$B$4)),"")',fmt)
    d.cell(last+2,1,"※ 実績は「スマホ」シートで入力します（この列は自動）。").font=Font(name=F)
    m.cell(MROW-1,1,f"■ 毎月の数字を入れる：{metric_name}（{unit}）").font=Font(name=F,bold=True)
    for col,w in zip("ABCDE",[10,18,18,18,22]): d.column_dimensions[col].width=w
    b=wb.create_sheet("ボード"); b["A1"]=f"現在地ボード：{title.split('：')[-1]}"; b["A1"].font=Font(name=F,bold=True,size=14)
    rng=f"記録!C2:C{last}"
    rng_a=f"記録!A2:A{last}"
    G='(前提!B3-前提!B2)'
    rows=[
     (3,"現在地",None,None,True),
     (4,f"最後に実績を入れた{period_word}",f'=IF(COUNT({rng})=0,"",LOOKUP(2,1/({rng}<>""),{rng_a}))',"0",False),
     (5,f"現在の{metric_name}（{unit}）",f'=IF(B4="","",LOOKUP(2,1/({rng}<>""),{rng}))',fmt,True),
     (6,f"目標（{unit}）","=前提!B3",fmt,False),
     (7,"目標までの進み具合",f'=IF(OR(B4="",{G}=0),"",(B5-前提!B2)/{G})',"0%",True),
     (8,"当初の計画と比べて",f'=IF(B4="","",IF({G}=0,"前提を確認（最初の値と目標が同じ）",IF((B5-INDEX(記録!B2:B{last},B4))*SIGN({G})>=0,"計画どおり、または先行","計画より遅れ")))',None,True),
     (9,f"残りの{period_word}数",'=IF(B4="","",前提!B4-B4)',"0",False),
     (10,f"目標に届くために必要な1か月あたりの変化（{unit}）",f'=IF(OR(B4="",{G}=0),"",IF((前提!B3-B5)*SIGN({G})<=0,0,IF(B9<=0,"",(前提!B3-B5)/B9)))',fmt,True),
     (11,f"これまでの1か月あたりの平均の変化（{unit}）",f'=IF(COUNT({rng})<2,"",(B5-前提!B2)/B4)',fmt,False),
     (12,"今のペースで目標に届く見込み（過去の数値を直線で延ばした参考計算）",f'=IF(B4="","",IF({G}=0,"前提を確認（最初の値と目標が同じ）",IF((前提!B3-B5)*SIGN({G})<=0,"達成済み",IF(B9<=0,"期限を過ぎています（前提の見直しが必要）",IF(COUNT({rng})<2,"記録が2か月分以上で表示",IF(B11*SIGN({G})<=0,"今のペースでは届かない（計画の見直しが必要）",IF(B4+(前提!B3-B5)/B11<=前提!B4,"期限内（"&ROUNDUP(B4+(前提!B3-B5)/B11,0)&"か月目ごろ）","期限より遅れる（"&ROUNDUP(B4+(前提!B3-B5)/B11,0)&"か月目ごろ）")))))))',None,True),
     (13,"入力のチェック",f'=IF(COUNTIF({rng},"<0")>0,"マイナスの値があります。確認してください",IF(B4="","実績を入力してください",IF(COUNT({rng})<B4,"入力していない月があります（現在地は最後に入力した月で計算）","OK")))',None,False),
    ]
    for r,k,f,f2,key in rows:
        if f is None:
            b.cell(r,1,k).font=Font(name=F,bold=True); continue
        lab(b,r,1,k); calc(b.cell(r,2),f,f2,key)
    b["A15"]="これからの予定（自動更新）"; b["A15"].font=Font(name=F,bold=True)
    b["A16"]=f"「記録」シートの「更新後の予定」列に、残りの{period_word}ごとに目指す値が自動で入ります。実績を入れるたびに作り直されます。"; b["A16"].font=Font(name=F)
    b.column_dimensions["A"].width=46; b.column_dimensions["B"].width=40
    b.add_chart(plan_chart(d,m,last-1,"当初の計画・実績・更新後の予定",unit,False),"F3")
    m["A2"]="■ 現在地"; m["A2"].font=Font(name=F,bold=True)
    for rr,(lb,src,f2) in enumerate([("現在地",'=ボード!B5',fmt),("目標",'=ボード!B6',fmt),("進み具合",'=ボード!B7',"0%"),("計画と比べて",'=ボード!B8',None),("必要なペース（1か月あたり）",'=ボード!B10',fmt),("見込み（参考計算）",'=ボード!B12',None),("入力のチェック",'=ボード!B13',None)],3):
        mres(m,rr,lb,src,f2,key=lb in ("現在地","進み具合","見込み（参考計算）"))
    m["A10"]=f"単位：{unit}。見込みは過去の数値を直線で延ばした参考計算です。"; m["A10"].font=Font(name=F,size=9)
    mchart(m,d,"計画・実績・予定",unit,2,last,[2,3,4]); fix_heights(m)
    if extra: extra(wb,b,d,last)
    return wb

def board1():
    def extra(wb,b,d,last):
        b["A18"]="AI化した業務"; b["A18"].font=Font(name=F,bold=True)
        lab(b,19,1,"AI化する予定の業務の数"); inp(b.cell(19,2),3,"0")
        lab(b,20,1,"AI化が終わった業務の数（スマホのシートから自動）"); calc(b.cell(20,2),'=スマホ!B'+str(MROW+last),"0")
        m=wb["スマホ"]; lab(m,MROW+last-1+1,1,"AI化が終わった業務の数"); inp(m.cell(MROW+last,2),2,"0")
        mres(m,MROW+last+1,"AI化の進み具合",'=ボード!B21',"0%")
        lab(b,21,1,"AI化の進み具合"); calc(b.cell(21,2),'=IF(OR(B19="",B19<=0),"",MIN(B20/B19,1))',"0%",True)
    return linear("① 現在地ボード：AI業務改善",
      ["毎月1回、「スマホ」シートの下にある「毎月の数字を入れる」欄に、先月の事務作業の合計時間（時間/月）を入れてください。",
       "入れると「ボード」シートで、現在地・目標までの進み具合・必要なペース・期限内に届く見込みが自動で計算され、残りの月の予定も作り直されます。",
       "「前提」シートの値は初回の計画書から転記します（月次更新プランでは運営者が見直します）。"],
      "時間/月",22,10,6,"月",[22,20,16],"事務作業時間",extra=extra)

def board3():
    def extra(wb,b,d,last):
        b["A18"]="段階（マイルストーン）の見込み"; b["A18"].font=Font(name=F,bold=True)
        header(b,19,["段階","到達の目安（累計時間）","状態","今のペースで到達する見込み"])
        ms=[("M1 商業簿記の基礎",60),("M2 工業簿記",130),("M3 過去問演習",200),("ゴール 受験",250)]
        for i,(nm,th) in enumerate(ms,20):
            inp(b.cell(i,1),nm); inp(b.cell(i,2),th,"0")
            calc(b.cell(i,3),f'=IF(B$5="","",IF(B$5>=B{i},"達成","未達成"))',None,True)
            calc(b.cell(i,4),f'=IFERROR(IF(B$5>=B{i},"—",IF(B$11<=0,"見込みなし","{"}"+""+"{"}"&ROUNDUP(B$4+(B{i}-B$5)/B$11,0)&"か月目ごろ")),"")',None,True)
        for i in range(20,24):
            b.cell(i,4).value=f'=IFERROR(IF(B$5="","",IF(B$5>=B{i},"—",IF(B$11="","記録が2か月分以上で表示",IF(B$11<=0,"見込みなし",ROUNDUP(B$4+(B{i}-B$5)/B$11,0)&"か月目ごろ")))),"")'
        m=wb["スマホ"]; r0=MROW+last+1
        m.cell(r0,1,"■ 段階ごとの見込み").font=Font(name=F,bold=True)
        for i in range(4):
            mres(m,r0+1+i,f'=ボード!A{20+i}',f'=IF(ボード!C{20+i}="達成","達成",ボード!D{20+i})',None,False)
        b["A25"]="※ 架空の記入例（簿記2級、合計250時間と仮定）。段階と目安は初回の計画書から転記します。"; b["A25"].font=Font(name=F)
        for col,w in zip("CD",[14,26]): b.column_dimensions[col].width=w
    return linear("③ 現在地ボード：目標達成ロードマップ",
      ["毎月1回、「スマホ」シートの下にある「毎月の数字を入れる」欄に、その月までの累計の取り組み時間（勉強時間など）を入れてください。",
       "入れると、現在地・進み具合・必要なペース・期限内に届く見込みが自動で計算され、残りの月の予定と、各段階に到達する見込みの時期が作り直されます。",
       "「前提」と「段階」は初回の計画書から転記します。"],
      "時間",0,250,7,"月",[30,55,80],"累計の取り組み時間",fmt="0",extra=extra)

def board2():
    wb=Workbook()
    usage(wb,"② 現在地ボード：家計・積立",
      ["毎月1回、「スマホ」シートの下にある「毎月の数字を入れる」欄に、月末の積立残高（評価額のおおよその額）を入れてください。",
       "入れると、現在地・目標までの進み具合と、期限までに目標に届くために必要な毎月の積立額が自動で計算され、残りの月の予定が作り直されます。",
       "利回りは仮定です（予測・保証ではありません）。0%やマイナスでも試してください。税金（課税口座は利益に原則20.315%）は含みません。銘柄や商品の選び方の助言は含みません。"])
    p=wb.create_sheet("前提"); p["A1"]="前提（仮定条件）"; p["A1"].font=Font(name=F,bold=True)
    items=[("最初の積立残高（円）",0,"#,##0"),("目標の積立残高（円）",1000000,"#,##0"),("期限までの月数",30,"0"),("仮定の利回り（年）",0.03,"0.0%"),("運用手数料（年）",0.001,"0.0%")]
    for i,(k,v,f) in enumerate(items,2): lab(p,i,1,k); inp(p.cell(i,2),v,f)
    lab(p,7,1,"月あたりの実質利回り（自動）"); calc(p.cell(7,2),"=(B5-B6)/12","0.000%")
    lab(p,8,1,"当初の計画の毎月の積立額（自動・円）"); calc(p.cell(8,2),'=IF(B7=0,(B3-B2)/B4,PMT(B7,B4,B2,-B3))',"#,##0")
    p.column_dimensions["A"].width=40; p.column_dimensions["B"].width=14
    n=30; last=n+1
    m=mobile(wb,"② スマホで見る：家計・積立")
    d=wb.create_sheet("記録")
    header(d,1,["月","当初の計画（円・仮定）","実績（円）※スマホのシートから自動","更新後の予定（円・仮定）"])
    act=[30500,61000,90000,118000]
    for k in range(1,n+1):
        r=k+1; calc(d.cell(r,1),k,"0")
        calc(d.cell(r,2),f'=IF(前提!$B$7=0,前提!$B$2+前提!$B$8*A{r},FV(前提!$B$7,A{r},-前提!$B$8,-前提!$B$2))',"#,##0")
        mr=MROW+k-1
        lab(m,mr,1,f"{k}か月目"); inp(m.cell(mr,2),act[k-1] if k-1<len(act) else None,"#,##0")
        calc(d.cell(r,3),f'=IF(スマホ!B{mr}="","",スマホ!B{mr})',"#,##0")
        calc(d.cell(r,4),f'=IFERROR(IF(OR(ボード!$B$4="",ボード!$B$10="",A{r}<=ボード!$B$4),"",IF(前提!$B$7=0,ボード!$B$5+ボード!$B$10*(A{r}-ボード!$B$4),FV(前提!$B$7,A{r}-ボード!$B$4,-ボード!$B$10,-ボード!$B$5))),"")',"#,##0")
    d.cell(last+2,1,"※ 実績は「スマホ」シートで入力します（この列は自動）。").font=Font(name=F)
    m.cell(MROW-1,1,"■ 毎月の数字を入れる：月末の積立残高（円）").font=Font(name=F,bold=True)
    for col,w in zip("ABCD",[8,22,18,22]): d.column_dimensions[col].width=w
    b=wb.create_sheet("ボード"); b["A1"]="現在地ボード：家計・積立（仮定条件にもとづく試算）"; b["A1"].font=Font(name=F,bold=True,size=14)
    rng=f"記録!C2:C{last}"
    rng_a=f"記録!A2:A{last}"
    rows=[(3,"現在地",None,None,False),
     (4,"最後に実績を入れた月",f'=IF(COUNT({rng})=0,"",LOOKUP(2,1/({rng}<>""),{rng_a}))',"0",False),
     (5,"現在の積立残高（円）",f'=IF(B4="","",LOOKUP(2,1/({rng}<>""),{rng}))',"#,##0",True),
     (6,"目標（円）","=前提!B3","#,##0",False),
     (7,"目標までの進み具合",'=IF(OR(B4="",前提!B3=前提!B2),"",(B5-前提!B2)/(前提!B3-前提!B2))',"0%",True),
     (8,"当初の計画と比べて",f'=IF(B4="","",IF(B5>=INDEX(記録!B2:B{last},B4),"計画どおり、または先行","計画より遅れ"))',None,True),
     (9,"残りの月数",'=IF(B4="","",前提!B4-B4)',"0",False),
     (10,"目標に届くために必要な毎月の積立額（更新後・円）",'=IF(B4="","",IF(B5>=前提!B3,0,IF(B9<=0,"",IF(前提!B7=0,(前提!B3-B5)/B9,PMT(前提!B7,B9,B5,-前提!B3)))))',"#,##0",True),
     (11,"当初の計画の毎月の積立額（円）","=前提!B8","#,##0",False),
     (12,"当初の積立額を続けた場合の期限時点の残高（円・仮定）",'=IF(B4="","",IF(B9<=0,B5,IF(前提!B7=0,B5+前提!B8*B9,FV(前提!B7,B9,-前提!B8,-B5))))',"#,##0",True),
     (13,"入力のチェック",f'=IF(COUNTIF({rng},"<0")>0,"マイナスの値があります。確認してください",IF(B4="","実績を入力してください",IF(B9<=0,"期限に達しました（前提の見直しが必要）",IF(COUNT({rng})<B4,"入力していない月があります（現在地は最後に入力した月で計算）","OK"))))',None,False)]
    for r,k,f,f2,key in rows:
        if f is None: b.cell(r,1,k).font=Font(name=F,bold=True); continue
        lab(b,r,1,k); calc(b.cell(r,2),f,f2,key)
    b["A15"]="計算の前提：月あたりの利回り＝（年の利回り−年の運用手数料）÷12。最初の残高は1か月目の初めにあり、毎月の積立は月末に入る。複利は月ごと。税金・物価上昇は含まない。"; b["A15"].font=Font(name=F)
    b["A16"]="※ 仮定の利回りによる機械的な計算で、予測・保証ではありません。投資には元本割れのリスクがあります。"; b["A16"].font=Font(name=F)
    b.column_dimensions["A"].width=48; b.column_dimensions["B"].width=18
    b.add_chart(plan_chart(d,m,last-1,"当初の計画・実績・更新後の予定（仮定）","円",False),"D3")
    m["A2"]="■ 現在地"; m["A2"].font=Font(name=F,bold=True)
    for rr,(lb,src,f2) in enumerate([("現在の積立残高",'=ボード!B5',"#,##0"),("目標",'=ボード!B6',"#,##0"),("進み具合",'=ボード!B7',"0%"),("計画と比べて",'=ボード!B8',None),("目標に必要な毎月の積立額（更新後）",'=ボード!B10',"#,##0"),("今の積立額を続けた場合の期限時点（仮定）",'=ボード!B12',"#,##0"),("入力のチェック",'=ボード!B13',None)],3):
        mres(m,rr,lb,src,f2,key=lb in ("現在の積立残高","進み具合","目標に必要な毎月の積立額（更新後）"))
    m["A10"]="単位：円。仮定の利回りによる計算で、予測・保証ではありません。元本割れのリスクがあります。税金は含みません。"; m["A10"].font=Font(name=F,size=9); m["A10"].alignment=Alignment(wrap_text=True)
    mchart(m,d,"計画・実績・予定（仮定）","円",2,last,[2,3,4]); fix_heights(m)
    return wb

def board4():
    wb=Workbook()
    usage(wb,"④ 現在地ボード：受験ロードマップ",
      ["模試・テストの結果が出たら、「スマホ」シートの下にある入力欄に、教科ごとの偏差値（または点数）を入れてください（保護者の方と一緒に）。",
       "入れると、教科ごとの現在地・志望の目安との差・受験までの模試1回あたりに必要な伸び・次の模試での目標が自動で計算されます。",
       "入試日程・制度は学校や教育委員会の公式情報で確認してください。合格を保証するものではありません。受験生の本名・学校名は書かないでください。"])
    subj=["英語","数学","国語","理科","社会"]
    m=mobile(wb,"④ スマホで見る：受験")
    d=wb.create_sheet("模試の記録")
    header(d,1,["回"]+subj+["模試の名前（会社・シリーズ）"])
    # スマホの入力欄：表（行＝回、列＝教科）。空欄は本当の空欄のままなので、グラフに0として出ない
    TG=21   # 志望の目安の行。1回目〜8回目は22〜29行
    m.cell(19,1,"■ 数字を入れる（模試ごと・偏差値または点数）").font=Font(name=F,bold=True)
    header(m,20,["回"]+["英","数","国","理","社"])
    lab(m,TG,1,"志望の目安",True)
    for j in range(5):
        inp(m.cell(TG,j+2),60,"0")
    calc(d.cell(2,1),"志望の目安")
    for j in range(5):
        calc(d.cell(2,j+2),f'=IF(スマホ!{"BCDEF"[j]}{TG}="","",スマホ!{"BCDEF"[j]}{TG})',"0")
    rows=[("1回目",[58,49,56,52,55]),("2回目",[59,52,56,54,55])]+[(f"{i}回目",[None]*5) for i in range(3,9)]
    m.cell(31,1,"■ 模試の名前（会社・シリーズ）").font=Font(name=F,bold=True)
    for i,(nm,vals) in enumerate(rows,3):
        mr=TG+i-2
        lab(m,mr,1,nm)
        for j,v in enumerate(vals):
            inp(m.cell(mr,j+2),v,"0")
            col="BCDEF"[j]
            calc(d.cell(i,j+2),f'=IF(スマホ!{col}{mr}="","",スマホ!{col}{mr})',"0")
        calc(d.cell(i,1),nm)
        nr=31+i-2
        lab(m,nr,1,nm); m.merge_cells(start_row=nr,start_column=2,end_row=nr,end_column=6)
        inp(m.cell(nr,2),"（例）◯◯模試" if vals[0] is not None else None)
        calc(d.cell(i,7),f'=IF(スマホ!B{nr}="","",スマホ!B{nr})')
    lab(d,12,1,"受験までに受ける模試の回数（全部で）"); calc(d.cell(12,2),"=スマホ!B17","0")
    lab(m,17,1,"受験までの模試の回数（全部）"); inp(m.cell(17,2),8,"0")
    d.cell(13,1,"※ 架空の記入例（中学2年・高校受験）。").font=Font(name=F)
    d.column_dimensions["A"].width=34
    for col in "BCDEF": d.column_dimensions[col].width=9
    b=wb.create_sheet("ボード"); b["A1"]="現在地ボード：受験"; b["A1"].font=Font(name=F,bold=True,size=14)
    lab(b,2,1,"受けた模試の回数"); calc(b.cell(2,2),"=MAX(COUNT(模試の記録!B3:B10),COUNT(模試の記録!C3:C10),COUNT(模試の記録!D3:D10),COUNT(模試の記録!E3:E10),COUNT(模試の記録!F3:F10))","0")
    lab(b,3,1,"受験までの残りの模試の回数"); calc(b.cell(3,2),"=MAX(模試の記録!B12-B2,0)","0")
    header(b,5,["教科","現在地","志望の目安","差","1回目からの平均の伸び（1回あたり・3回以上で表示）","目安に届くために必要な伸び（1回あたり）","次の模試の目標（参考）","今のペースの場合（参考計算）"])
    for k,s in enumerate(subj):
        r=6+k; col="BCDEF"[k]; rng=f"模試の記録!{col}3:{col}10"
        lab(b,r,1,s)
        calc(b.cell(r,2),f'=IFERROR(LOOKUP(2,1/({rng}<>""),{rng}),"")',"0",True)
        calc(b.cell(r,3),f"=模試の記録!{col}2","0")
        calc(b.cell(r,4),f'=IFERROR(B{r}-C{r},"")',"0;-0;0")
        calc(b.cell(r,5),f'=IFERROR(IF(COUNT({rng})<3,"",(B{r}-模試の記録!{col}3)/(COUNT({rng})-1)),"")',"0.0")
        calc(b.cell(r,6),f'=IFERROR(IF(D{r}>=0,0,IF($B$3=0,"",-D{r}/$B$3)),"")',"0.0",True)
        calc(b.cell(r,7),f'=IFERROR(IF(D{r}>=0,B{r},ROUND(B{r}+F{r},0)),"")',"0",True)
        calc(b.cell(r,8),f'=IFERROR(IF(B{r}="","",IF(D{r}>=0,"目安に到達",IF($B$3=0,"残りの模試がありません",IF(E{r}="","予測データ不足（3回以上で表示）",IF(E{r}>=F{r},"このペースなら届く計算","このペースでは届かない計算（配分の見直し）"))))),"")',None,True)
    b["A12"]="※ 見込みと次の模試の目標は、過去の数値を直線で延ばした参考計算で、学力や合格の可能性の予測ではありません。時間配分の見直しは、初回の計画書の方針（または模試後の見直し）を優先してください。"; b["A12"].font=Font(name=F)
    b["A13"]="※ 偏差値は、模試の実施会社・受験者層・問題の難しさで変わります。できるだけ同じ会社・同じシリーズの模試の結果を記録してください。違う模試の偏差値は、単純な伸びとして比べられません。"; b["A13"].font=Font(name=F)
    b.column_dimensions["A"].width=30
    for col,w in zip("BCDEFGH",[10,10,8,16,18,12,30]): b.column_dimensions[col].width=w
    def exam_chart(small):
        ch=LineChart(); ch.title="教科別の推移"; ch.legend.position="b"
        ch.height,ch.width=(7,9.5) if small else (9,20)
        for j,sb in enumerate(subj):
            sr=Series(Reference(m,min_col=j+2,min_row=22,max_row=29),title=sb); sr.smooth=False; ch.series.append(sr)
        ch.set_categories(Reference(m,min_col=1,min_row=22,max_row=29))
        return ch
    b.add_chart(exam_chart(False),"A15")
    m["A2"]="■ 現在地（最新の模試）"; m["A2"].font=Font(name=F,bold=True)
    m.column_dimensions["A"].width=11
    for col in "BCDEF": m.column_dimensions[col].width=6.5
    m.merge_cells("A1:F1")
    for k,sb in enumerate(subj):
        r=6+k
        mres(m,3+k,f"{sb}",f'=IF(ボード!B{r}="","",ボード!B{r}&"（目安"&ボード!C{r}&"・差"&TEXT(ボード!D{r},"+0;-0;0")&"）")',None,False)
        m.merge_cells(start_row=3+k,start_column=2,end_row=3+k,end_column=6)
    m.cell(9,1,"■ 次の模試の目標（参考）").font=Font(name=F,bold=True)
    for k,sb in enumerate(subj):
        r=6+k
        mres(m,10+k,sb,f'=IF(ボード!G{r}="","",ボード!G{r}&"　"&ボード!H{r})',None,True)
        m.merge_cells(start_row=10+k,start_column=2,end_row=10+k,end_column=6); m.row_dimensions[10+k].height=30
    m["A15"]="見込みと目標は過去の数値を直線で延ばした参考計算で、合格の可能性の予測ではありません。できるだけ同じ会社・同じシリーズの模試で比べてください。"; m["A15"].font=Font(name=F,size=9); m["A15"].alignment=Alignment(wrap_text=True,vertical="top")
    m.merge_cells("A15:F15"); m.row_dimensions[15].height=48
    m.add_chart(exam_chart(True),"A42")
    m["A41"]="■ 教科別の推移"; m["A41"].font=Font(name=F,bold=True)
    return wb

for name,fn in [("board-01-business.xlsx",board1),("board-02-money.xlsx",board2),("board-03-goal.xlsx",board3),("board-04-exam.xlsx",board4)]:
    wb=fn()
    from openpyxl.worksheet.datavalidation import DataValidation
    for ws in wb.worksheets:
        dv=DataValidation(type="decimal",operator="greaterThanOrEqual",formula1="0",allow_blank=True,showErrorMessage=True,errorTitle="入力の確認",error="0以上の数字を入れてください")
        cells=[c.coordinate for row in ws.iter_rows() for c in row if c.fill==YEL and (isinstance(c.value,(int,float)) or c.value is None) and not (ws.title=="前提" and c.row in (5,6) and name.startswith("board-02"))]
        if cells:
            ws.add_data_validation(dv)
            for co in cells: dv.add(co)
    wb.save(f"{OUT}/{name}")
print("ok")

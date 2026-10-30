"""現在地ボードの境界値テスト。

使い方: python3 test_boards.py <recalc.py のパス>
各ボードのコピーに境界値を入れ、LibreOffice で再計算し、
数式エラーが0であることと、表示が期待どおりであることを確認する。
"""
import json
import shutil
import subprocess
import sys
import tempfile
from pathlib import Path

from openpyxl import load_workbook

HERE = Path(__file__).resolve().parent
RECALC = sys.argv[1]


MROW = 27  # build_boards.py と同じ。数字の入力欄はスマホシートにある


def to_input(sheet, cell):
    """記録シートの実績セルなどを、実際の入力欄（スマホシート）の位置に読み替える。"""
    col, row = cell[0], int(cell[1:])
    if sheet == "記録" and col == "C":
        return "スマホ", f"B{MROW + row - 2}"
    if sheet == "模試の記録":
        if col == "B" and row == 12:
            return "スマホ", "B17"
        return "スマホ", f"{col}{row + 19}"  # 志望の目安＝21行、1回目〜8回目＝22〜29行
    return sheet, cell


def run(board, edits, checks, name):
    tmp = Path(tempfile.mkdtemp()) / board
    shutil.copy(HERE / board, tmp)
    wb = load_workbook(tmp)
    for sheet, cell, value in edits:
        sheet, cell = to_input(sheet, cell)
        wb[sheet][cell] = value
    wb.save(tmp)
    out = json.loads(subprocess.run([sys.executable, RECALC, str(tmp), "90"], capture_output=True, text=True).stdout)
    ok = out.get("status") == "success" and out.get("total_errors") == 0
    ws = load_workbook(tmp, data_only=True)
    fails = []
    for sheet, cell, expect in checks:
        got = ws[sheet][cell].value
        good = expect(got) if callable(expect) else got == expect
        if not good:
            fails.append(f"{sheet}!{cell}={got!r}")
    status = "PASS" if ok and not fails else "FAIL"
    print(f"[{status}] {board} / {name}" + ("" if ok else f" errors={out}") + (f" {fails}" if fails else ""))
    return status == "PASS"


def clear(sheet, col, rows):
    return [(sheet, f"{col}{r}", None) for r in rows]


def has(text):
    return lambda v: isinstance(v, str) and text in v


results = []
B1 = "board-01-business.xlsx"
results += [
    run(B1, [], [("ボード", "B12", has("期限内")), ("ボード", "B13", "OK")], "記入例（減らす目標）"),
    run(B1, clear("記録", "C", range(2, 8)), [("ボード", "B12", None), ("ボード", "B13", has("実績を入力"))], "実績が未入力"),
    run(B1, clear("記録", "C", range(3, 8)), [("ボード", "B12", has("2か月分以上"))], "記録が1か月だけ"),
    run(B1, [("前提", "B3", 22)], [("ボード", "B7", None), ("ボード", "B12", has("前提を確認"))], "最初の値と目標が同じ（分母0）"),
    run(B1, [("記録", "C5", 8)], [("ボード", "B12", "達成済み"), ("ボード", "B10", 0)], "目標を超えて達成"),
    run(B1, [("記録", "C5", 15), ("記録", "C6", 14), ("記録", "C7", 13)], [("ボード", "B12", has("期限を過ぎて"))], "期限を過ぎて未達"),
    run(B1, [("記録", "C3", 24), ("記録", "C4", 25)], [("ボード", "B12", has("届かない"))], "悪化している"),
    run(B1, [("記録", "C3", None)], [("ボード", "B4", 3), ("ボード", "B13", has("入力していない月"))], "途中の月が空欄"),
    run(B1, [("記録", "C4", -5)], [("ボード", "B13", has("マイナス"))], "マイナスの値"),
]
B2 = "board-02-money.xlsx"
results += [
    run(B2, [], [("ボード", "B8", "計画より遅れ"), ("ボード", "B10", lambda v: isinstance(v, float) and 32000 < v < 33500), ("ボード", "B13", "OK")], "記入例"),
    run(B2, [("前提", "B5", 0), ("前提", "B6", 0)], [("ボード", "B10", lambda v: isinstance(v, (int, float)) and abs(v - (1000000 - 118000) / 26) < 1)], "利回り0%"),
    run(B2, [("前提", "B5", -0.02)], [("ボード", "B10", lambda v: isinstance(v, float) and v > 0)], "マイナスの利回り"),
    run(B2, [("記録", "C5", 1200000)], [("ボード", "B10", 0)], "目標を超えて達成"),
    run(B2, [("前提", "B3", 0)], [("ボード", "B7", None)], "最初の値と目標が同じ（分母0）"),
    run(B2, clear("記録", "C", range(2, 32)), [("ボード", "B13", has("実績を入力"))], "実績が未入力"),
    run(B2, [("前提", "B4", 4)], [("ボード", "B10", None), ("ボード", "B13", has("期限に達しました"))], "期限に達した"),
]
B3 = "board-03-goal.xlsx"
results += [
    run(B3, [], [("ボード", "B12", has("期限より遅れる")), ("ボード", "C21", "未達成"), ("ボード", "D21", has("か月目ごろ"))], "記入例"),
    run(B3, clear("記録", "C", range(3, 9)), [("ボード", "D21", has("2か月分以上"))], "記録が1か月だけ"),
    run(B3, [("記録", "C5", 260)], [("ボード", "B12", "達成済み"), ("ボード", "C23", "達成")], "目標を超えて達成"),
    run(B3, [("前提", "B3", 0)], [("ボード", "B12", has("前提を確認"))], "最初の値と目標が同じ（分母0）"),
]
B4 = "board-04-exam.xlsx"
results += [
    run(B4, [], [("ボード", "E7", None), ("ボード", "H7", has("予測データ不足"))], "記録2回（データ不足）"),
    run(B4, [("模試の記録", c + "5", v) for c, v in zip("BCDEF", [60, 55, 57, 55, 56])],
        [("ボード", "H6", "目安に到達"), ("ボード", "H7", has("計算"))], "記録3回"),
    run(B4, [("模試の記録", "B12", 2)], [("ボード", "B3", 0), ("ボード", "H7", has("残りの模試がありません"))], "残りの模試0回"),
    run(B4, clear("模試の記録", "B", range(3, 11)) + clear("模試の記録", "C", range(3, 11)) + clear("模試の記録", "D", range(3, 11)) + clear("模試の記録", "E", range(3, 11)) + clear("模試の記録", "F", range(3, 11)),
        [("ボード", "B6", None), ("ボード", "H6", None)], "結果が未入力"),
]
print(f"{sum(results)}/{len(results)} passed")
sys.exit(0 if all(results) else 1)
